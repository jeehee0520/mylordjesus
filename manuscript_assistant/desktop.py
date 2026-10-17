"""Desktop entry point: runs the app on this computer and opens it in the browser.

Used both from source (`python desktop.py`) and as the frozen executable built by
PyInstaller (see packaging/). The server listens on 127.0.0.1 only, so nothing
is reachable from other computers and no manuscript leaves the machine.

    ManuscriptAssistant              start and open the browser
    ManuscriptAssistant --selftest   check that every component works, then exit (used in CI)
"""
from __future__ import annotations

import os
import socket
import sys
import threading
import time
import urllib.request
import webbrowser
from pathlib import Path

APP_NAME = "Manuscript Assistant"


def base_dir() -> Path:
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))


def free_port(start: int = 8501, tries: int = 50) -> int:
    for port in range(start, start + tries):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            try:
                s.bind(("127.0.0.1", port))
                return port
            except OSError:
                continue
    raise RuntimeError("사용할 수 있는 포트가 없습니다.")


def wait_ready(url: str, timeout: float = 120) -> bool:
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            with urllib.request.urlopen(f"{url}/_stcore/health", timeout=2) as resp:
                if resp.status == 200:
                    return True
        except OSError:
            pass
        time.sleep(0.5)
    return False


def check_components() -> list[str]:
    """Exercise every bundled dependency on synthetic data (no files are read or written)."""
    import io

    import docx
    from PIL import Image

    from core import guideline_loader
    from core.figure_processor import FigureOptions, process
    from core.manuscript_formatter import format_manuscript
    from core.profile_extractor import extract_profile
    from core.proofreader import proofread
    from core.template_builder import build_template
    from core.track_changes import apply_corrections

    done = []
    g = guideline_loader.from_text(
        "Thank you for your interest in Self Test Journal. Use A4 paper, Arial 11-point font, double-spaced. "
        "Please provide a structured abstract of no more than 200 words. Provide 4-6 keywords.", "selftest")
    profile = extract_profile(g)
    assert profile.get("abstract.max_words") == 200
    done.append("규정 분석")
    template = build_template(profile)
    done.append("템플릿 생성")
    d = docx.Document()
    d.add_heading("Introduction", 1)
    d.add_paragraph("This is is a testt  sentence.")
    buf = io.BytesIO()
    d.save(buf)
    res = format_manuscript(buf.getvalue(), profile)
    assert res.guard.ok
    done.append("서식 변환")
    issues = [i for i in proofread(buf.getvalue(), profile) if i.fixable]
    assert issues and apply_corrections(buf.getvalue(), issues).ok
    done.append("오류 점검·변경 추적")
    img = io.BytesIO()
    Image.new("RGB", (600, 300), "white").save(img, "PNG", dpi=(300, 300))
    fig = process("t.png", img.getvalue(), FigureOptions(dpi=300, width_mm=50))
    assert fig.outputs and not fig.error
    pdf = next(o.data for o in fig.outputs if o.filename.endswith(".pdf"))
    assert process("t.pdf", pdf, FigureOptions(dpi=150, width_mm=50)).outputs
    guideline_loader.load(template, "template.docx")
    done.append("Figure 변환 (PNG·PDF)")
    try:
        guideline_loader.load(pdf, "t.pdf")
    except ValueError:
        pass  # the test PDF has no text; this only proves the PDF reader is bundled
    done.append("PDF 읽기")
    return done


def setup_output() -> None:
    """A windowed app (macOS .app) has no console: write messages to a log file instead.

    With a console, never crash on characters the console encoding cannot show.
    """
    if sys.stdout is None or sys.stderr is None:
        log_dir = Path.home() / ".manuscript_assistant"
        log_dir.mkdir(exist_ok=True)
        log = open(log_dir / "app.log", "w", encoding="utf-8", buffering=1)
        sys.stdout = sys.stdout or log
        sys.stderr = sys.stderr or log
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass


def main(argv: list[str]) -> int:
    setup_output()
    selftest = "--selftest" in argv
    base = base_dir()
    os.chdir(base)  # Streamlit reads .streamlit/config.toml from here
    if str(base) not in sys.path:
        sys.path.insert(0, str(base))

    if selftest:
        try:
            for item in check_components():
                print(f"OK  {item}", flush=True)
        except Exception as exc:  # pragma: no cover - reported in CI
            import traceback
            traceback.print_exc()
            print(f"FAIL  {exc!r}")
            return 1

    port = free_port()
    url = f"http://localhost:{port}"
    os.environ["MA_DESKTOP"] = "1"

    def watch():
        if not wait_ready(url):
            print("앱을 시작하지 못했습니다.", flush=True)
            os._exit(1)
        if selftest:
            with urllib.request.urlopen(url, timeout=10) as resp:
                ok = resp.status == 200
            print("OK  웹 화면" if ok else "FAIL  웹 화면", flush=True)
            os._exit(0 if ok else 1)
        print(f"\n{APP_NAME}: {url}", flush=True)
        print("브라우저가 열리지 않으면 위 주소를 입력하세요. 끝낼 때는 앱 왼쪽의 '프로그램 종료'를 누르거나 이 창을 닫으세요.\n")
        webbrowser.open(url)

    threading.Thread(target=watch, daemon=True).start()

    from streamlit.web import bootstrap

    flags = {
        "server.port": port,
        "server.address": "127.0.0.1",
        "server.headless": True,
        "global.developmentMode": False,
        "browser.gatherUsageStats": False,
        "server.fileWatcherType": "none",
    }
    bootstrap.load_config_options(flag_options=flags)
    bootstrap.run(str(base / "app.py"), False, [], flags)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
