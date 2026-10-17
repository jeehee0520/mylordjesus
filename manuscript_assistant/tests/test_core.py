import io
import zipfile

import docx
import pytest
from PIL import Image
from pypdf import PdfReader

from core import content_guard, guideline_loader
from core.figure_processor import FigureOptions, process
from core.manuscript_formatter import FormatOptions, format_manuscript
from core.ooxml import DocxPackage, w
from core.profile import DEFAULT, FOUND, Profile
from core.profile_extractor import extract_profile
from core.proofreader import ProofOptions, proofread
from core.template_builder import build_template
from core.textmodel import DocView
from core.track_changes import AUTHOR, apply_corrections


@pytest.fixture
def profile(guideline_text):
    p = extract_profile(guideline_loader.from_text(guideline_text, "test.txt"))
    p.confirmed = True
    return p


# ---------------------------------------------------------------- guidelines
def test_extracts_explicit_rules_with_evidence(profile):
    expected = {
        "journal_name": "Journal of Synthetic Testing",
        "page.size": "A4", "page.margin_cm": 2.5, "text.font": "Arial", "text.size_pt": 11.0,
        "text.line_spacing": 2.0, "text.line_numbers": True, "text.page_numbers": True,
        "abstract.max_words": 200, "abstract.structured": True,
        "abstract.headings": ["Background", "Methods", "Results", "Conclusions"],
        "keywords.min": 4, "keywords.max": 6, "style.spelling": "American", "refs.style": "Vancouver",
        "refs.citation": "numbers in brackets", "refs.order_of_appearance": True,
        "tables.footnote_symbols": ["†", "‡", "§"], "figures.separate_files": True,
        "figures.dpi_halftone": 300, "figures.dpi_line": 1200, "figures.width_single_mm": 85.0,
        "figures.width_double_mm": 175.0, "figures.formats": ["TIFF", "EPS", "PDF"],
        "statements.data_availability": True, "statements.conflict_of_interest": True,
    }
    for key, value in expected.items():
        rule = profile.rules[key]
        assert rule.value == value, key
        assert rule.status == FOUND, key
        assert rule.quote and rule.location, key
    assert profile.rules["structure.order"].value == [
        "Title page", "Abstract", "Keywords", "Introduction", "Methods", "Results", "Discussion",
        "Acknowledgments", "References", "Tables", "Figure legends"]


def test_missing_rules_are_not_invented(profile):
    assert profile.rules["figures.dpi_combination"].status == DEFAULT
    assert profile.rules["length.max_words"].status == "not_specified"
    assert profile.rules["statements.trial_registration"].value is None


def test_profile_json_roundtrip(profile):
    again = Profile.from_json(profile.to_json())
    assert again.rules["abstract.headings"].value == profile.rules["abstract.headings"].value
    assert again.confirmed


def test_loads_docx_and_pdf_guidelines(guideline_text):
    d = docx.Document()
    for line in guideline_text.splitlines():
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    g = guideline_loader.load(buf.getvalue(), "g.docx")
    assert extract_profile(g).get("abstract.max_words") == 200
    with pytest.raises(ValueError):
        guideline_loader.load(b"x", "g.doc")


# ---------------------------------------------------------------- template
def test_template_has_profile_formatting(profile):
    data = build_template(profile)
    d = docx.Document(io.BytesIO(data))
    sec = d.sections[0]
    assert round(sec.left_margin.cm, 2) == 2.5
    assert round(sec.page_width.cm, 1) == 21.0
    assert sec._sectPr.find(w("lnNumType")) is not None
    assert d.styles["Normal"].font.name == "Arial"
    assert d.styles["Normal"].font.size.pt == 11
    assert d.styles["Normal"].paragraph_format.line_spacing == 2.0
    headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert headings.index("Introduction") < headings.index("Acknowledgments") < headings.index("References")
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Background:" in text and "4–6 keywords" in text
    assert b"PAGE" in zipfile.ZipFile(io.BytesIO(data)).read("word/footer1.xml")


# ---------------------------------------------------------------- formatting
def test_formatting_changes_only_formatting(manuscript, profile):
    res = format_manuscript(manuscript, profile, FormatOptions())
    assert res.guard.ok, res.guard.details
    pkg = DocxPackage.from_bytes(res.data)
    doc = pkg.xml("word/document.xml")
    mar = doc.find(f".//{w('pgMar')}")
    assert mar.get(w("left")) == str(round(2.5 / 2.54 * 1440))
    assert doc.find(f".//{w('lnNumType')}") is not None
    runs = [r for r in doc.iter(w("r")) if r.find(w("t")) is not None and not any(a.tag == w("tc") for a in r.iterancestors())]
    assert all(r.find(f"{w('rPr')}/{w('rFonts')}").get(w("ascii")) == "Arial" for r in runs)
    assert all(r.find(f"{w('rPr')}/{w('sz')}").get(w("val")) == "22" for r in runs)
    # existing revisions and the citation field survive
    assert any(e.get(w("author")) == "Reviewer" for e in doc.iter(w("ins")))
    assert "EN.CITE" in "".join(t.text for t in doc.iter(w("instrText")))
    original = content_guard.document_signature(DocxPackage.from_bytes(manuscript))
    assert content_guard.compare(original, content_guard.document_signature(pkg)).ok


def test_guard_detects_any_text_change(manuscript):
    pkg = DocxPackage.from_bytes(manuscript)
    before = content_guard.document_signature(pkg)
    t = next(pkg.xml("word/document.xml").iter(w("t")))
    t.text = t.text + "x"
    after = content_guard.document_signature(DocxPackage.from_bytes(pkg.to_bytes()))
    res = content_guard.compare(before, after)
    assert not res.ok and res.details


def test_guard_ignores_run_splitting(manuscript):
    pkg = DocxPackage.from_bytes(manuscript)
    before = content_guard.document_signature(pkg)
    import copy
    t = next(t for t in pkg.xml("word/document.xml").iter(w("t")) if len(t.text or "") > 4)
    run = t.getparent()
    clone = copy.deepcopy(run)
    clone.find(w("t")).text = t.text[2:]
    t.text = t.text[:2]
    run.addnext(clone)
    assert content_guard.compare(before, content_guard.document_signature(DocxPackage.from_bytes(pkg.to_bytes()))).ok


def test_compliance_report(manuscript, profile):
    res = format_manuscript(manuscript, profile)
    items = {f.item: f for f in res.findings}
    assert items["키워드 개수"].level == "fail"          # 3 keywords, 4-6 required
    assert items["키워드 구분자"].level == "warn"
    assert "Data availability statement" in items and items["Data availability statement"].level == "fail"


# ---------------------------------------------------------------- proofreading
def test_proofreader_finds_typical_errors(manuscript, profile):
    issues = proofread(manuscript, profile, ProofOptions())
    found = {(i.category, i.original, i.suggestion) for i in issues}
    assert ("띄어쓰기", "  ", " ") in found
    assert any(i.category == "반복" and i.suggestion == "" for i in issues)
    assert ("저널 철자 규칙", "colour", "color") in found      # American spelling journal
    assert ("저널 철자 규칙", "analysed", "analyzed") in found
    assert ("철자", "wrod", "word") in found                  # inside another author's insertion
    assert any(i.category == "표기 일관성" and i.original == "Bland-Altman" for i in issues)
    assert any(i.category == "그림 번호" for i in issues)       # Figure 2 cited before Figure 1
    assert any(i.category == "띄어쓰기" and i.message.startswith("문장부호 앞") for i in issues)
    # the double space before a citation field is editable, the field itself is not
    field_issue = next(i for i in issues if i.context.startswith("Cited"))
    assert field_issue.fixable


def test_corrections_are_tracked_and_verified(manuscript, profile):
    issues = [i for i in proofread(manuscript, profile, ProofOptions()) if i.fixable]
    res = apply_corrections(manuscript, issues)
    assert res.ok, res.checks
    assert len(res.applied) == len(issues)
    doc = DocxPackage.from_bytes(res.tracked).xml("word/document.xml")
    ours = [e for e in doc.iter(w("ins"), w("del")) if e.get(w("author")) == AUTHOR]
    assert ours
    clean = DocView.load(DocxPackage.from_bytes(res.clean))
    text = "\n".join(p.text for p in clean.paras)
    assert "inserted word here" in text and "the the" not in text and "analyzed" in text
    # other author's revisions remain untouched in both files
    clean_doc = DocxPackage.from_bytes(res.clean).xml("word/document.xml")
    assert any(e.get(w("author")) == "Reviewer" for e in clean_doc.iter(w("del")))


def test_locked_field_is_never_edited(manuscript, profile):
    from core.proofreader import Issue
    view = DocView.load(DocxPackage.from_bytes(manuscript))
    p = next(p for p in view.paras if p.text.startswith("Cited"))
    pos = p.text.index("2")
    bogus = Issue(p.index, pos, pos + 1, "2", "3", "test", "", "높음", fixable=True, id="x")
    res = apply_corrections(manuscript, [bogus])
    assert res.applied == [] and res.skipped


def test_report_only_issues_are_not_applied(manuscript, profile):
    issues = [i for i in proofread(manuscript, profile, ProofOptions()) if not i.fixable]
    res = apply_corrections(manuscript, issues)
    assert res.ok and res.applied == []


# ---------------------------------------------------------------- figures
def _png(size=(1000, 500), dpi=300, mode="RGB"):
    buf = io.BytesIO()
    Image.new(mode, size, "white").save(buf, "PNG", dpi=(dpi, dpi))
    return buf.getvalue()


def test_raster_resize_to_width_and_dpi():
    res = process("fig.png", _png((2000, 1000), 300), FigureOptions(dpi=300, width_mm=85))
    assert not res.error and not res.warnings
    png = next(o for o in res.outputs if o.filename == "fig.png")
    im = Image.open(io.BytesIO(png.data))
    assert im.size == (1004, 502) and round(im.info["dpi"][0]) == 300
    pdf = next(o for o in res.outputs if o.filename == "fig.pdf")
    box = PdfReader(io.BytesIO(pdf.data)).pages[0].mediabox
    assert abs(float(box.width) / 72 * 25.4 - 85) < 0.1
    assert pdf.measured["dpi"] == 300


def test_large_upscale_needs_consent():
    small = _png((300, 150), 72)
    res = process("small.png", small, FigureOptions(dpi=600, width_mm=170))
    assert res.error and not res.outputs
    res = process("small.png", small, FigureOptions(dpi=600, width_mm=170, allow_large_upscale=True))
    assert res.outputs and any("업스케일" in w for w in res.warnings)


def test_pdf_input_stays_vector_and_scales():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(400, 200))
    c.line(10, 10, 390, 190)
    c.save()
    res = process("vec.pdf", buf.getvalue(), FigureOptions(dpi=600, width_mm=85))
    pdf = next(o for o in res.outputs if o.filename.endswith(".pdf"))
    assert pdf.measured["픽셀"] == "벡터" and abs(pdf.measured["폭(mm)"] - 85) < 0.2
    png = next(o for o in res.outputs if o.filename.endswith(".png"))
    assert Image.open(io.BytesIO(png.data)).size[0] == round(85 / 25.4 * 600)


def test_cmyk_pdf_and_rgb_png():
    res = process("c.png", _png(), FigureOptions(dpi=300, color_mode="CMYK"))
    pdf = next(o for o in res.outputs if o.filename.endswith(".pdf"))
    png = next(o for o in res.outputs if o.filename.endswith(".png"))
    assert pdf.measured["컬러"] == "DeviceCMYK" and png.measured["컬러"] == "RGB"
