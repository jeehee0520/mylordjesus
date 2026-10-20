import io

import docx
import pytest
from pypdf import PdfReader

from core import guideline_loader
from core.figure_processor import FigureOptions, _find_soffice, parse_pages, process
from core.ooxml import DocxPackage
from core.profile_extractor import extract_profile
from core.table_footnotes import plan_table_footnotes
from core.textmodel import DocView
from core.track_changes import apply_corrections


@pytest.fixture
def profile(guideline_text):
    return extract_profile(guideline_loader.from_text(guideline_text, "g.txt"))


def _sup(paragraph, text):
    run = paragraph.add_run(text)
    run.font.superscript = True


def table_doc(cell_markers, notes) -> bytes:
    d = docx.Document()
    d.add_heading("Tables", 1)
    d.add_paragraph("Table 1. Results")
    t = d.add_table(rows=1, cols=len(cell_markers))
    for cell, (label, marker) in zip(t.rows[0].cells, cell_markers):
        p = cell.paragraphs[0]
        p.add_run(label)
        _sup(p, marker)
    for marker, text in notes:
        p = d.add_paragraph()
        _sup(p, marker)
        p.add_run(text)
    d.add_heading("Figure legends", 1)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _plan(data, profile):
    return plan_table_footnotes(DocView.load(DocxPackage.from_bytes(data)), profile)


def test_numbered_notes_become_journal_symbols(profile):
    data = table_doc([("Age", "1"), ("CRP", "2"), ("ESR", "2")],
                     [("1", "Measured at baseline."), ("2", "P < 0.05 versus control.")])
    issues, plans = _plan(data, profile)
    assert plans[0].mapping == {"1": "†", "2": "*"}  # general note → †, P value note → *
    res = apply_corrections(data, issues)
    assert res.ok, res.checks
    text = "\n".join(p.text for p in DocView.load(DocxPackage.from_bytes(res.clean)).paras)
    assert "Age†" in text and "CRP*" in text and "*P < 0.05" in text
    again, _ = _plan(res.clean, profile)
    assert again == []  # idempotent


def test_symbols_used_for_the_wrong_purpose_are_swapped(profile):
    data = table_doc([("N/A", "*"), ("0.3", "†")],
                     [("*", "Not estimable."), ("†", "P-value adjusted with Hochberg adjustment.")])
    issues, plans = _plan(data, profile)
    assert plans[0].mapping == {"*": "†", "†": "*"}
    assert len(issues) == 4 and apply_corrections(data, issues).ok


def test_order_follows_first_use_in_the_table(profile):
    data = table_doc([("A", "b"), ("B", "a")], [("a", "First note."), ("b", "Second note.")])
    _, plans = _plan(data, profile)
    assert plans[0].mapping == {"b": "†", "a": "‡"}


def test_superscripts_without_a_footnote_are_left_alone(profile):
    data = table_doc([("BMI, kg/m", "2"), ("Age", "1")], [("1", "Median.")])
    issues, plans = _plan(data, profile)
    assert plans[0].mapping == {"1": "†"}
    assert all(i.original == "1" for i in issues)


# ---------------------------------------------------------------- slides / PDF pages
def test_parse_pages():
    assert parse_pages("") is None
    assert parse_pages("1, 3-5,3") == [1, 3, 4, 5]
    with pytest.raises(ValueError):
        parse_pages("1;x")


def _two_page_pdf():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(720, 405))  # 16:9 slide
    for _ in range(2):
        c.rect(200, 150, 200, 100, fill=1)  # content only in the middle
        c.showPage()
    c.save()
    return buf.getvalue()


def test_trim_and_page_selection_keep_vector():
    res = process("deck.pdf", _two_page_pdf(), FigureOptions(dpi=300, width_mm=85, trim_whitespace=True, pages=[2]))
    assert [o.filename for o in res.outputs] == ["deck_p2.pdf", "deck_p2.png"]
    pdf = res.outputs[0]
    box = PdfReader(io.BytesIO(pdf.data)).pages[0].mediabox
    assert abs(float(box.width) / 72 * 25.4 - 85) < 0.2
    assert float(box.height) / float(box.width) < 0.6  # cropped to the 2:1 rectangle, not 16:9 page
    assert pdf.measured["픽셀"] == "벡터"


@pytest.mark.skipif(_find_soffice() is None, reason="LibreOffice not installed")
def test_pptx_slides_become_figures():
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for i in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(2), Inches(2), Inches(4), Inches(1))
        tb.text_frame.text = f"Panel {i + 1}"
    buf = io.BytesIO()
    prs.save(buf)
    res = process("slides.pptx", buf.getvalue(), FigureOptions(dpi=300, width_mm=80, trim_whitespace=True))
    assert not res.error, res.error
    assert sorted(o.filename for o in res.outputs) == ["slides_slide1.pdf", "slides_slide1.png", "slides_slide2.pdf", "slides_slide2.png"]


# ---------------------------------------------------------------- proofreading tab
def test_proofreader_includes_table_footnotes(profile):
    from core.proofreader import ProofOptions, proofread
    data = table_doc([("Age", "1")], [("1", "P < 0.05.")])
    issues = [i for i in proofread(data, profile, ProofOptions()) if i.category == "표 각주 기호"]
    assert {(i.original, i.suggestion) for i in issues} == {("1", "*")}
    assert all(i.default_selected for i in issues)
    assert not [i for i in proofread(data, profile, ProofOptions(table_footnotes=False)) if i.category == "표 각주 기호"]


# ---------------------------------------------------------------- saved guidelines
def test_pasted_guidelines_are_saved_and_reloaded(tmp_path, monkeypatch, guideline_text, profile):
    from core import storage
    from core.profile import Profile
    monkeypatch.setenv("MA_DATA_DIR", str(tmp_path))
    first = storage.save(guideline_text)
    again = storage.save(guideline_text + "\n")  # same text, different whitespace: one entry
    assert first.id == again.id and len(storage.list_saved()) == 1
    profile.confirmed = True
    storage.save(guideline_text, name="Journal of Synthetic Testing", profile_json=profile.to_json())
    entry = storage.list_saved()[0]
    assert entry.name == "Journal of Synthetic Testing" and "확정됨" in entry.label
    restored = Profile.from_json(entry.profile_json)
    assert restored.confirmed and restored.get("abstract.max_words") == 200
    storage.save(guideline_text)  # a later autosave keeps the confirmed profile
    assert Profile.from_json(storage.load(first.id).profile_json).confirmed
    storage.delete(first.id)
    assert storage.list_saved() == []
