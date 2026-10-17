"""Synthetic fixtures. No real manuscript or guideline text is stored in the repository.

Set MA_GUIDELINE and MA_MANUSCRIPT to local files to also run the tests in
test_real_files.py against your own documents.
"""
from __future__ import annotations

import io

import docx
import pytest
from docx.oxml import parse_xml
from docx.shared import Pt

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

GUIDELINE_TEXT = """Instructions to Authors
Thank you for your interest in Journal of Synthetic Testing.
Manuscripts must be prepared on A4 paper with margins of 2.5 cm, using Arial 11-point font,
double-spaced throughout, with continuous line numbering. Pages should be numbered.
The manuscript should be presented in the following order:
i. Title page;
ii. Abstract and keywords;
iii. Main text - Introduction, Methods, Results, and Discussion;
iv. Acknowledgments;
v. References;
vi. Tables;
vii. Figure legends.
Please provide a structured abstract of no more than 200 words with the headings: Background/Methods/Results/Conclusions.
Provide 4-6 keywords. The journal uses American spelling.
References should be numbered consecutively in order of appearance and cited with numbers in square brackets, following the Vancouver style.
Footnote symbols: †, ‡, § should be used and * reserved for P-values.
Figures should be supplied as separate files. Photographs should be at least 300 dpi, line art 1200 dpi.
Single-column figures are 85 mm wide and double-column figures 175 mm wide. Accepted figure formats are TIFF, EPS and PDF.
Authors must include a data availability statement and a conflict of interest statement.
"""


def make_manuscript() -> bytes:
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading("Title page", 1)
    d.add_paragraph("Title: A study of the the TNF-alpha pathway")
    d.add_heading("Abstract", 1)
    d.add_paragraph("Background: We analysed  data. Methods: x. Results: y. Conclusions: z.")
    d.add_paragraph("Keywords: alpha; beta, gamma")
    d.add_heading("Introduction", 1)
    p = d.add_paragraph("Inflammatory bowel disease (IBD) is common")
    sup = p.add_run("1")
    sup.font.superscript = True
    p.add_run(". The colour of the IBD tumour was analysed (Figure 2) and (Figure 1).")
    d.add_heading("Methods", 1)
    d.add_paragraph("Samples were randomised at a single centre .")
    d.add_heading("Results", 1)
    d.add_paragraph("Results are summarised in Table 1. Bland-Altman and Bland–Altman and Bland–Altman plots.")
    d.add_heading("Discussion", 1)
    d.add_paragraph("We discuss.")
    d.add_heading("References", 1)
    d.add_paragraph("1. Smith A. A paper. J Test 2020;1:1-2.")
    d.add_paragraph("2. Jones B. Another paper. J Test 2021;2:3-4.")
    d.add_heading("Tables", 1)
    d.add_paragraph("Table 1. Baseline characteristics")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Age"
    t.cell(0, 1).text = "50"
    d.add_heading("Figure legends", 1)
    d.add_paragraph("Figure 1. A figure.")
    d.add_paragraph("Figure 2. Another figure.")

    body = d.element.body
    # a paragraph with an existing tracked insertion and deletion by another author
    tracked = parse_xml(
        f'<w:p {W}><w:r><w:t xml:space="preserve">Existing </w:t></w:r>'
        f'<w:ins w:id="901" w:author="Reviewer" w:date="2024-01-01T00:00:00Z"><w:r><w:t xml:space="preserve">inserted wrod here</w:t></w:r></w:ins>'
        f'<w:del w:id="902" w:author="Reviewer" w:date="2024-01-01T00:00:00Z"><w:r><w:delText xml:space="preserve"> removed</w:delText></w:r></w:del>'
        f'<w:r><w:t xml:space="preserve"> text.</w:t></w:r></w:p>')
    # a citation field (EndNote-like) whose result must never be edited
    field = parse_xml(
        f'<w:p {W}><w:r><w:t xml:space="preserve">Cited  here</w:t></w:r>'
        f'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> ADDIN EN.CITE </w:instrText></w:r>'
        f'<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:t>2</w:t></w:r>'
        f'<w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t>.</w:t></w:r></w:p>')
    intro_heading = [p for p in body.iter("{*}p") if "".join(t.text or "" for t in p.iter("{*}t")) == "Methods"][0]
    intro_heading.addprevious(tracked)
    intro_heading.addprevious(field)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def manuscript() -> bytes:
    return make_manuscript()


@pytest.fixture
def guideline_text() -> str:
    return GUIDELINE_TEXT
